"""Export finančního přehledu analytiky (Excel / PDF)."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, timedelta
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.db.models import Q, Sum
from django.utils import timezone as dj_tz
from django.utils.translation import gettext_lazy as _

from core.money import format_castka, get_mena_symbol
from .analytika_data import (
    TrenerRateLookup,
    charges_sum_between,
    hraci_kredit_qs,
    month_end_from_start,
    ostatni_naklady_sum_for_month,
    payments_sum_between,
)
from .models import Hrac, OstatniNaklad, TrenerPlatba, Transakce, Trening

EXPORT_PERIODS = {
    "rok": _("Poslední rok"),
    "pulrok": _("Poslední půlrok"),
    "cela_doba": _("Celá doba"),
}


@dataclass
class ReportRow:
    section: str
    label: str
    value: str
    key: str = ""
    note: str = ""
    raw: float | int | None = None
    kind: str = "text"  # text | money | percent | hours | count


def period_bounds(period: str, today: date | None = None) -> tuple[date, date, str]:
    today = today or dj_tz.localdate()
    if period == "rok":
        return today - timedelta(days=365), today, EXPORT_PERIODS["rok"]
    if period == "pulrok":
        return today - timedelta(days=183), today, EXPORT_PERIODS["pulrok"]
    if period == "cela_doba":
        first_dt = Trening.objects.order_by("datum").values_list("datum", flat=True).first()
        if first_dt:
            d = dj_tz.localtime(first_dt).date() if dj_tz.is_aware(first_dt) else first_dt.date()
        else:
            d = today
        return d, today, EXPORT_PERIODS["cela_doba"]
    raise ValueError(f"Neznámé období: {period}")


def _money(v: Decimal | float | int) -> str:
    return format_castka(v, thousands=True, nbsp=True)


def _pct(v: Decimal) -> str:
    return f"{v.quantize(Decimal('0.1'), rounding=ROUND_HALF_UP)} %"


def _hours_str(mins: int) -> str:
    h = (Decimal(mins) / Decimal(60)).quantize(Decimal("0.01"))
    return f"{h} h"


def _other_costs_between(day_from: date, day_to: date) -> Decimal:
    return Decimal(
        OstatniNaklad.objects.filter(mesic__gte=day_from, mesic__lte=day_to).aggregate(s=Sum("castka"))["s"] or 0
    )


def _monthly_rows(day_from: date, day_to: date) -> list[dict]:
    month_start = day_from.replace(day=1)
    rows = []
    czech = {1: "Led", 2: "Úno", 3: "Bře", 4: "Dub", 5: "Kvě", 6: "Čer",
             7: "Čvc", 8: "Srp", 9: "Zář", 10: "Říj", 11: "Lis", 12: "Pro"}
    while month_start <= day_to:
        m_end = min(month_end_from_start(month_start), day_to)
        m_from = max(month_start, day_from)
        if m_from <= m_end:
            mins = Trening.objects.filter(
                datum__date__gte=m_from, datum__date__lte=m_end
            ).aggregate(s=Sum("delka_minut"))["s"] or 0
            charged = charges_sum_between(m_from, m_end)
            paid = payments_sum_between(m_from, m_end)
            rows.append({
                "mesic": f"{czech[month_start.month]} {month_start.year}",
                "hodiny": float((Decimal(mins) / Decimal(60)).quantize(Decimal("0.01"))),
                "nauctovano": float(charged),
                "platby": float(paid),
            })
        if month_start.month == 12:
            month_start = month_start.replace(year=month_start.year + 1, month=1, day=1)
        else:
            month_start = month_start.replace(month=month_start.month + 1, day=1)
    return rows


def build_financial_report(period: str, today: date | None = None) -> dict:
    today = today or dj_tz.localdate()
    day_from, day_to, period_label = period_bounds(period, today)
    if day_from > day_to:
        day_from = day_to

    trener_ids = list(
        Trening.objects.exclude(trener_id__isnull=True).values_list("trener_id", flat=True).distinct()
    )
    rate_lookup = TrenerRateLookup(trener_ids)

    training_qs = Trening.objects.filter(datum__date__gte=day_from, datum__date__lte=day_to)
    total_min = training_qs.aggregate(s=Sum("delka_minut"))["s"] or 0
    trainings_count = training_qs.count()

    charged = charges_sum_between(day_from, day_to)
    paid = payments_sum_between(day_from, day_to)
    charges_count = Transakce.objects.filter(
        typ=Transakce.Typ.NAUCTOVANO,
        trening__datum__date__gte=day_from,
        trening__datum__date__lte=day_to,
    ).count()
    payments_count = Transakce.objects.filter(
        vytvoreno__date__gte=day_from,
        vytvoreno__date__lte=day_to,
        typ__in=[Transakce.Typ.PLATBA, Transakce.Typ.VRATKA],
    ).count()

    trener_earned = rate_lookup.earned_from_rows(
        training_qs.values("trener_id", "datum", "delka_minut")
    )
    trener_paid = Decimal(
        TrenerPlatba.objects.filter(
            vytvoreno__date__gte=day_from,
            vytvoreno__date__lte=day_to,
        ).aggregate(s=Sum("castka"))["s"] or 0
    )
    other_costs = _other_costs_between(day_from, day_to)
    unpaid = charged - paid
    cashflow = paid - trener_paid - other_costs
    gross_profit = charged - trener_earned
    net_profit = cashflow
    hours_dec = (Decimal(total_min) / Decimal(60)).quantize(Decimal("0.01"))
    margin = (gross_profit / charged * 100) if charged > 0 else Decimal("0")
    revenue_per_hour = (charged / hours_dec) if hours_dec > 0 else Decimal("0")

    agregace = hraci_kredit_qs().aggregate(
        dluhy=Sum("_kredit_calculated", filter=Q(_kredit_calculated__lt=Decimal(0))),
        prebytky=Sum("_kredit_calculated", filter=Q(_kredit_calculated__gt=Decimal(0))),
        celkem=Sum("_kredit_calculated"),
    )
    total_debt = agregace.get("dluhy") or Decimal(0)
    total_surplus = agregace.get("prebytky") or Decimal(0)
    total_balance = agregace.get("celkem") or Decimal(0)

    total_trener_earned_all = rate_lookup.earned_from_rows(
        Trening.objects.values("trener_id", "datum", "delka_minut")
    )
    total_trener_paid_all = Decimal(TrenerPlatba.objects.aggregate(s=Sum("castka"))["s"] or 0)
    trener_to_pay = total_trener_earned_all - total_trener_paid_all

    rows: list[ReportRow] = [
        ReportRow("Aktivita", "Odehrané hodiny", _hours_str(total_min), "hours", raw=float(hours_dec), kind="hours"),
        ReportRow("Aktivita", "Počet tréninků", str(trainings_count), "trainings", raw=trainings_count, kind="count"),
        ReportRow("Finance", "Naúčtováno", _money(charged), "charged", raw=float(charged), kind="money"),
        ReportRow("Finance", "Přijaté platby", _money(paid), "paid", raw=float(paid), kind="money"),
        ReportRow("Finance", "Nezaplaceno (v období)", _money(unpaid), "unpaid", raw=float(unpaid), kind="money"),
        ReportRow("Finance", "Počet naúčtování", str(charges_count), "charges_count", raw=charges_count, kind="count"),
        ReportRow("Finance", "Počet plateb", str(payments_count), "payments_count", raw=payments_count, kind="count"),
        ReportRow("Účetnictví", "Náklady na trenéry (earned)", _money(trener_earned), "trener_earned", raw=float(trener_earned), kind="money"),
        ReportRow("Účetnictví", "Výplaty trenérům", _money(trener_paid), "trener_paid", raw=float(trener_paid), kind="money"),
        ReportRow("Účetnictví", "Ostatní náklady", _money(other_costs), "other_costs", raw=float(other_costs), kind="money"),
        ReportRow("Účetnictví", "Cash flow", _money(cashflow), "cashflow", raw=float(cashflow), kind="money"),
        ReportRow("Účetnictví", "Hrubý zisk", _money(gross_profit), "gross_profit", raw=float(gross_profit), kind="money"),
        ReportRow("Účetnictví", "Čistý zisk", _money(net_profit), "net_profit", raw=float(net_profit), kind="money"),
        ReportRow("Účetnictví", "Marže", _pct(margin), "margin", raw=float(margin), kind="percent"),
        ReportRow("Účetnictví", "Výnosy na hodinu", _money(revenue_per_hour), "revenue_per_hour", raw=float(revenue_per_hour), kind="money"),
        ReportRow("Stav k datu exportu", "K vyplacení trenérům", _money(trener_to_pay), "trener_balance",
                  note="Celkový stav, ne jen za vybrané období", raw=float(trener_to_pay), kind="money"),
        ReportRow("Stav k datu exportu", "Dluhy hráčů", _money(abs(total_debt)), "player_debt",
                  note="Aktuální bilance", raw=float(abs(total_debt)), kind="money"),
        ReportRow("Stav k datu exportu", "Přebytky hráčů", _money(total_surplus), "player_surplus",
                  note="Aktuální bilance", raw=float(total_surplus), kind="money"),
        ReportRow("Stav k datu exportu", "Finanční rovnováha hráčů", _money(total_balance), "player_balance",
                  note="Aktuální bilance", raw=float(total_balance), kind="money"),
    ]

    try:
        from core.models import SystemNastaveni
        club_name = SystemNastaveni.load().nazev_klubu or "TenisSystém"
    except Exception:
        club_name = "TenisSystém"

    return {
        "title": f"Finanční přehled – {club_name}",
        "period_slug": period,
        "period_label": period_label,
        "date_from": day_from,
        "date_to": day_to,
        "generated_at": dj_tz.localtime(),
        "rows": rows,
        "monthly": _monthly_rows(day_from, day_to),
        "players_count": Hrac.objects.count(),
        "club_name": club_name,
    }


def render_report_excel(report: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    # Barvy a styly
    brand = "E67817"
    brand_fill = PatternFill("solid", fgColor=brand)
    header_fill = PatternFill("solid", fgColor="F1F5F9")
    section_fill = PatternFill("solid", fgColor="FFF4E8")
    alt_fill = PatternFill("solid", fgColor="FAFAFA")
    white_font = Font(color="FFFFFF", bold=True, size=14)
    title_font = Font(bold=True, size=12, color="1A1D16")
    section_font = Font(bold=True, size=11, color=brand)
    header_font = Font(bold=True, size=10, color="1A1D16")
    normal_font = Font(size=10, color="1A1D16")
    muted_font = Font(size=9, color="64748B", italic=True)
    thin = Side(style="thin", color="E2E8F0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    money_fmt = f'#,##0 "{get_mena_symbol()}"'
    hours_fmt = '#,##0.0" h"'
    count_fmt = '#,##0'

    def apply_value(cell, row: ReportRow):
        if row.kind == "money" and row.raw is not None:
            cell.value = row.raw
            cell.number_format = money_fmt
        elif row.kind == "percent" and row.raw is not None:
            cell.value = row.raw / 100
            cell.number_format = '0.0%'
        elif row.kind == "hours" and row.raw is not None:
            cell.value = row.raw
            cell.number_format = hours_fmt
        elif row.kind == "count" and row.raw is not None:
            cell.value = row.raw
            cell.number_format = count_fmt
        else:
            cell.value = row.value

    # ── List 1: Přehled ──
    ws = wb.active
    ws.title = "Přehled"
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:D1")
    c = ws["A1"]
    c.value = report["title"]
    c.font = white_font
    c.fill = brand_fill
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 32

    meta_rows = [
        ("Období", report["period_label"]),
        ("Rozsah", f'{report["date_from"].strftime("%d.%m.%Y")} – {report["date_to"].strftime("%d.%m.%Y")}'),
        ("Vygenerováno", report["generated_at"].strftime("%d.%m.%Y %H:%M")),
        ("Počet hráčů", report["players_count"]),
    ]
    row_idx = 3
    for label, val in meta_rows:
        ws.cell(row=row_idx, column=1, value=label).font = Font(bold=True, size=10, color="64748B")
        ws.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=4)
        ws.cell(row=row_idx, column=2, value=val).font = normal_font
        row_idx += 1

    row_idx += 1
    current_section = None
    for item in report["rows"]:
        if item.section != current_section:
            current_section = item.section
            ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=4)
            sec_cell = ws.cell(row=row_idx, column=1, value=current_section)
            sec_cell.font = section_font
            sec_cell.fill = section_fill
            sec_cell.alignment = Alignment(horizontal="left", vertical="center")
            sec_cell.border = border
            for col in range(2, 5):
                ws.cell(row=row_idx, column=col).fill = section_fill
                ws.cell(row=row_idx, column=col).border = border
            ws.row_dimensions[row_idx].height = 22
            row_idx += 1

            for col, hdr in enumerate(["Ukazatel", "Hodnota", "Poznámka"], start=1):
                cell = ws.cell(row=row_idx, column=col, value=hdr if col < 3 else "")
                cell.font = header_font
                cell.fill = header_fill
                cell.border = border
                cell.alignment = Alignment(horizontal="left" if col == 1 else "right" if col == 2 else "left")
            ws.merge_cells(start_row=row_idx, start_column=3, end_row=row_idx, end_column=4)
            row_idx += 1

        data_row = row_idx
        ws.cell(row=data_row, column=1, value=item.label).font = normal_font
        ws.cell(row=data_row, column=1).border = border
        ws.cell(row=data_row, column=1).alignment = Alignment(vertical="center")

        val_cell = ws.cell(row=data_row, column=2)
        apply_value(val_cell, item)
        val_cell.font = Font(bold=True, size=10, color="1A1D16")
        val_cell.border = border
        val_cell.alignment = Alignment(horizontal="right", vertical="center")

        ws.merge_cells(start_row=data_row, start_column=3, end_row=data_row, end_column=4)
        note_cell = ws.cell(row=data_row, column=3, value=item.note or "")
        note_cell.font = muted_font
        note_cell.border = border
        note_cell.alignment = Alignment(wrap_text=True, vertical="center")

        if (data_row % 2) == 0:
            for col in range(1, 5):
                ws.cell(row=data_row, column=col).fill = alt_fill

        row_idx += 1

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 14
    ws.freeze_panes = "A3"

    # ── List 2: Měsíční přehled ──
    ws2 = wb.create_sheet("Měsíční přehled")
    ws2.sheet_view.showGridLines = False

    headers = ["Měsíc", "Hodiny", "Naúčtováno", "Platby", "Nezaplaceno"]
    ws2.merge_cells("A1:E1")
    ws2["A1"].value = f'Měsíční přehled – {report["period_label"]}'
    ws2["A1"].font = white_font
    ws2["A1"].fill = brand_fill
    ws2["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws2.row_dimensions[1].height = 28

    for col, hdr in enumerate(headers, start=1):
        cell = ws2.cell(row=3, column=col, value=hdr)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="right" if col > 1 else "left")
    ws2.cell(row=3, column=1).alignment = Alignment(horizontal="left")

    r = 4
    totals = {"hodiny": 0.0, "nauctovano": 0.0, "platby": 0.0, "nezaplaceno": 0.0}
    for m in report["monthly"]:
        nez = m["nauctovano"] - m["platby"]
        totals["hodiny"] += m["hodiny"]
        totals["nauctovano"] += m["nauctovano"]
        totals["platby"] += m["platby"]
        totals["nezaplaceno"] += nez

        ws2.cell(row=r, column=1, value=m["mesic"]).font = normal_font
        ws2.cell(row=r, column=2, value=m["hodiny"]).number_format = hours_fmt
        ws2.cell(row=r, column=3, value=m["nauctovano"]).number_format = money_fmt
        ws2.cell(row=r, column=4, value=m["platby"]).number_format = money_fmt
        ws2.cell(row=r, column=5, value=nez).number_format = money_fmt

        for col in range(1, 6):
            ws2.cell(row=r, column=col).border = border
            ws2.cell(row=r, column=col).alignment = Alignment(
                horizontal="right" if col > 1 else "left", vertical="center"
            )
            if r % 2 == 0:
                ws2.cell(row=r, column=col).fill = alt_fill
        r += 1

    # Součtový řádek
    ws2.cell(row=r, column=1, value="CELKEM").font = Font(bold=True, size=10)
    ws2.cell(row=r, column=2, value=totals["hodiny"]).number_format = hours_fmt
    ws2.cell(row=r, column=3, value=totals["nauctovano"]).number_format = money_fmt
    ws2.cell(row=r, column=4, value=totals["platby"]).number_format = money_fmt
    ws2.cell(row=r, column=5, value=totals["nezaplaceno"]).number_format = money_fmt
    for col in range(1, 6):
        cell = ws2.cell(row=r, column=col)
        cell.font = Font(bold=True, size=10, color="1A1D16")
        cell.fill = section_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="right" if col > 1 else "left", vertical="center")

    for col, width in enumerate([16, 12, 16, 16, 16], start=1):
        ws2.column_dimensions[get_column_letter(col)].width = width
    ws2.freeze_panes = "A4"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pdf_font_name() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        Path(settings.BASE_DIR) / "static" / "fonts" / "DejaVuSans.ttf",
        Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
        Path("/Library/Fonts/Arial Unicode.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]
    for path in candidates:
        if path.is_file():
            name = "TsCzechFont"
            if name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(name, str(path)))
            return name
    return "Helvetica"


def render_report_pdf(report: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font = _pdf_font_name()
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=report["title"],
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TsTitle",
        parent=styles["Heading1"],
        fontName=font,
        fontSize=16,
        spaceAfter=6,
    )
    sub_style = ParagraphStyle(
        "TsSub",
        parent=styles["Normal"],
        fontName=font,
        fontSize=10,
        textColor=colors.HexColor("#475569"),
        spaceAfter=12,
    )
    section_style = ParagraphStyle(
        "TsSection",
        parent=styles["Heading2"],
        fontName=font,
        fontSize=11,
        textColor=colors.HexColor("#E67817"),
        spaceBefore=10,
        spaceAfter=4,
    )
    cell_style = ParagraphStyle("TsCell", parent=styles["Normal"], fontName=font, fontSize=9)

    story = [
        Paragraph(report["title"], title_style),
        Paragraph(
            f'{report["period_label"]} · '
            f'{report["date_from"].strftime("%d.%m.%Y")} – {report["date_to"].strftime("%d.%m.%Y")} · '
            f'Vygenerováno {report["generated_at"].strftime("%d.%m.%Y %H:%M")}',
            sub_style,
        ),
    ]

    current_section = None
    table_data = [["Ukazatel", "Hodnota"]]
    section_tables: list[tuple[str, list]] = []

    def flush_section():
        nonlocal table_data, current_section
        if current_section and len(table_data) > 1:
            section_tables.append((current_section, table_data))
        table_data = [["Ukazatel", "Hodnota"]]

    for row in report["rows"]:
        if row.section != current_section:
            flush_section()
            current_section = row.section
        note = f' <font size="7" color="#64748b">({row.note})</font>' if row.note else ""
        table_data.append([
            Paragraph(row.label + note, cell_style),
            Paragraph(row.value.replace("\u00a0", " "), cell_style),
        ])
    flush_section()

    for section_name, data in section_tables:
        story.append(Paragraph(section_name, section_style))
        tbl = Table(data, colWidths=[105 * mm, 65 * mm])
        tbl.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), font),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5f9")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#1a1d16")),
            ("FONTNAME", (0, 0), (-1, 0), font),
            ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.HexColor("#cbd5e1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#fafafa")]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        story.append(tbl)
        story.append(Spacer(1, 4 * mm))

    if report["monthly"]:
        story.append(Paragraph("Měsíční přehled", section_style))
        m_data = [["Měsíc", "Hodiny", "Naúčtováno", "Platby"]]
        for m in report["monthly"]:
            m_data.append([
                m["mesic"],
                f'{m["hodiny"]:.1f} h',
                _money(m["nauctovano"]).replace("\u00a0", " "),
                _money(m["platby"]).replace("\u00a0", " "),
            ])
        m_tbl = Table(m_data, colWidths=[40 * mm, 28 * mm, 48 * mm, 48 * mm])
        m_tbl.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), font),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5f9")),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#e2e8f0")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#fafafa")]),
            ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
        ]))
        story.append(m_tbl)

    doc.build(story)
    return buf.getvalue()


def export_filename(report: dict, ext: str) -> str:
    slug = report["period_slug"]
    stamp = report["generated_at"].strftime("%Y%m%d")
    club = (report.get("club_name") or "tenissystem").strip().lower()
    club_slug = "".join(ch if ch.isalnum() else "-" for ch in club).strip("-") or "tenissystem"
    return f"analytika-{club_slug}_{slug}_{stamp}.{ext}"
